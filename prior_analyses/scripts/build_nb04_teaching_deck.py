"""Build a styled teaching deck (PowerPoint) for notebook 04.

Generates docs/notebook_4_teaching_deck.pptx — same visual style as
scripts/build_teaching_deck.py (notebooks 1-3) — explaining the faithfulness
audit step by step: what each check is, why it was run, and what it found.

Every number is read from the committed results JSONs
(docs/faithfulness_probe_results.json, docs/posthoc_audit_results.json,
docs/roadmap_followup_results.json), which match the re-executed notebook 04.

Usage:
    python scripts/build_nb04_teaching_deck.py
"""

from __future__ import annotations

import json
import tempfile
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from scipy.stats import norm

REPO_ROOT = Path(__file__).resolve().parents[1]
DOCS = REPO_ROOT / "docs"
OUT_PPTX = DOCS / "notebook_4_teaching_deck.pptx"
CHART_DIR = Path(tempfile.mkdtemp(prefix="nb04_charts_"))

probe = json.loads((DOCS / "faithfulness_probe_results.json").read_text())
audit = json.loads((DOCS / "posthoc_audit_results.json").read_text())
roadmap = json.loads((DOCS / "roadmap_followup_results.json").read_text())

N_TEST = probe["dataset"]["direct_chart_by_split"]["test"]  # 276
SD = float(np.sqrt(0.25 / N_TEST))
MDE2 = audit["analytic_power_min_detectable_accuracy"]["n_276"]["two_sided_alpha_05"]
ABL = probe["single_split_ablation"]
CV = probe["repeated_cv"]
FW = audit["family_wise_permutation"]
TF = probe["tfidf_text_probe"]
BIG = roadmap["expanded_baseline_all_pairs"]
CAL = roadmap["calibration"]
SIM = roadmap["similarity_eda"]
PLB = roadmap["paper_label_balance"]

# Palette (identical to the notebooks 1-3 deck)
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
ORANGE = RGBColor(0xE0, 0x7B, 0x39)
DARK = RGBColor(0x26, 0x26, 0x26)
GRAY_BG = RGBColor(0xF0, 0xF2, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
M_NAVY, M_ORANGE, M_GRAY = "#1F3A5F", "#E07B39", "#9AA5B1"


def save_fig(fig, name):
    p = CHART_DIR / name
    fig.savefig(p, dpi=200, bbox_inches="tight")
    plt.close(fig)
    return p


# ------------------------------------------------------------------ charts

def chart_chance():
    xs = np.linspace(0.38, 0.66, 400)
    fig, ax = plt.subplots(figsize=(9, 4.2))
    ax.plot(xs, norm.pdf(xs, 0.5, SD), color=M_NAVY, label=f"pure guessing, n={N_TEST} (SD {SD:.3f})")
    ax.axvspan(0.5 - 1.96 * SD, 0.5 + 1.96 * SD, color=M_NAVY, alpha=0.10,
               label=f"95% chance band: {0.5 - 1.96 * SD:.1%} – {0.5 + 1.96 * SD:.1%}")
    ax.axvline(0.5, color="gray", lw=1)
    ax.axvline(ABL["full_notebook_stack"]["accuracy"], color=M_ORANGE, ls="--",
               label=f"headline (full stack): {ABL['full_notebook_stack']['accuracy']:.1%}")
    ax.axvline(MDE2, color="seagreen", ls=":",
               label=f"true skill needed for reliable detection: {MDE2:.1%}")
    ax.set_xlabel("test accuracy")
    ax.set_yticks([])
    ax.legend(loc="upper left", fontsize=9)
    ax.spines[["top", "right", "left"]].set_visible(False)
    return save_fig(fig, "chance.png")


def chart_null_ranges():
    rows = [("TF-IDF claim text", TF["direct_chart"])] + [
        (k, ABL[k]) for k in ("claim_vec", "pair_text_vec", "image_vec", "claim+image", "full_notebook_stack")
    ]
    fig, ax = plt.subplots(figsize=(9, 4.2))
    for i, (name, r) in enumerate(rows):
        ypos = len(rows) - 1 - i
        ax.plot([r["null_min"], r["null_max"]], [ypos, ypos], color=M_NAVY, lw=2, alpha=0.35)
        ax.plot([r["null_p05"], r["null_p95"]], [ypos, ypos], color=M_NAVY, lw=7, alpha=0.75)
        ax.plot(r["null_mean"], ypos, "o", color="black", ms=5)
        ax.plot(r["accuracy"], ypos, "x", color=M_ORANGE, ms=10, mew=3)
    ax.axvline(0.5, color="gray", lw=1)
    ax.set_yticks(range(len(rows)))
    ax.set_yticklabels([n for n, _ in reversed(rows)], fontsize=9)
    ax.set_xlabel("test accuracy")
    handles = [
        plt.Line2D([], [], color=M_NAVY, lw=7, alpha=0.75, label="scrambled-label models: 5th–95th pct"),
        plt.Line2D([], [], color=M_NAVY, lw=2, alpha=0.35, label="scrambled-label models: full range (300 runs)"),
        plt.Line2D([], [], color="black", marker="o", ls="", label="scrambled-label mean"),
        plt.Line2D([], [], color=M_ORANGE, marker="x", ls="", mew=3, label="real model's score"),
    ]
    ax.legend(handles=handles, loc="upper left", fontsize=8)
    ax.spines[["top", "right"]].set_visible(False)
    return save_fig(fig, "null_ranges.png")


def chart_ablation_cv():
    names = ["claim_vec", "image_vec", "claim+image"]
    labels = ["claim text only", "image only", "claim + image"]
    single = [ABL[n]["accuracy"] for n in names]
    cv_mean = [CV[n]["cv_balanced_accuracy_mean"] for n in names]
    cv_std = [CV[n]["cv_balanced_accuracy_std"] for n in names]
    x = np.arange(len(names))
    w = 0.35
    fig, ax = plt.subplots(figsize=(9, 4.2))
    ax.bar(x - w / 2, single, w, color=M_NAVY, label="official split (train on 140)")
    ax.bar(x + w / 2, cv_mean, w, yerr=cv_std, capsize=5, color=M_ORANGE,
           label="repeated CV, 50 splits (train on ~333)")
    for i, (s, c) in enumerate(zip(single, cv_mean)):
        ax.text(i - w / 2, s + 0.004, f"{s:.3f}", ha="center", fontsize=9)
        ax.text(i + w / 2, c + cv_std[i] + 0.004, f"{c:.3f}", ha="center", fontsize=9)
    ax.axhline(0.5, color="gray", ls="--")
    ax.set_xticks(x)
    ax.set_xticklabels(labels)
    ax.set_ylim(0.40, 0.64)
    ax.set_ylabel("accuracy")
    ax.legend(fontsize=9)
    ax.spines[["top", "right"]].set_visible(False)
    return save_fig(fig, "ablation_cv.png")


def chart_six_mirror():
    fig, ax = plt.subplots(figsize=(9, 3.0))
    for ypos, (label, d, color) in enumerate([
        ("worst of six tries", FW["null_min"], M_NAVY),
        ("best of six tries", FW["null_max"], M_ORANGE),
    ]):
        ax.plot([d["min"], d["max"]], [ypos, ypos], color=color, lw=2, alpha=0.4)
        ax.plot([d["p05"], d["p95"]], [ypos, ypos], color=color, lw=9, alpha=0.75)
        ax.plot(d["mean"], ypos, "o", color="black", ms=6)
        ax.annotate(f"mean {d['mean']:.1%}", (d["mean"], ypos + 0.2), ha="center", fontsize=9)
    ax.axvline(0.5, color="gray", lw=1)
    ax.plot(FW["observed_best_accuracy"], 1, "x", color="black", ms=11, mew=3)
    ax.annotate(f"observed best: {FW['observed_best_accuracy']:.1%}",
                (FW["observed_best_accuracy"], 1.28), ha="center", fontsize=9, fontweight="bold")
    ax.set_yticks([0, 1])
    ax.set_yticklabels(["worst of six tries", "best of six tries"])
    ax.set_ylim(-0.5, 1.75)
    ax.set_xlabel("test accuracy of scrambled-label models")
    ax.spines[["top", "right"]].set_visible(False)
    return save_fig(fig, "six_mirror.png")


def chart_bigpop():
    keys = ["claim_vec", "image_vec", "claim+image", "full_notebook_stack", "metadata_onehots"]
    labels = ["claim text", "image", "claim + image", "full stack", "metadata\none-hots"]
    vals = [BIG[k]["accuracy"] for k in keys]
    fig, ax = plt.subplots(figsize=(9, 4.0))
    bars = ax.bar(labels, vals, color=[M_NAVY, M_ORANGE, M_NAVY, M_NAVY, M_GRAY])
    for b, v in zip(bars, vals):
        ax.text(b.get_x() + b.get_width() / 2, v + 0.003, f"{v:.3f}", ha="center", fontsize=10)
    ax.axhline(0.5, color="gray", ls="--", label="guessing (50%)")
    ax.set_ylim(0.45, 0.56)
    ax.set_ylabel("test accuracy (train 504 → test 996)")
    ax.legend(fontsize=9)
    ax.spines[["top", "right"]].set_visible(False)
    return save_fig(fig, "bigpop.png")


def chart_calibration():
    fig, axes = plt.subplots(1, 2, figsize=(9, 3.6))
    axes[0].bar(["this model", "always answer\n\"50/50 unsure\""],
                [CAL["brier_score"], CAL["brier_score_always_0.5"]], color=[M_ORANGE, M_GRAY], width=0.5)
    axes[0].text(0, CAL["brier_score"] + 0.008, f"{CAL['brier_score']}", ha="center", fontsize=11, fontweight="bold")
    axes[0].text(1, CAL["brier_score_always_0.5"] + 0.008, f"{CAL['brier_score_always_0.5']}", ha="center", fontsize=11)
    axes[0].set_ylabel("Brier score (lower = better)")
    axes[0].set_title("probability quality", fontsize=11)
    frac = CAL["frac_predictions_in_0.4_0.6"]
    axes[1].bar(["near-certain\n(outside 0.4–0.6)", "honest \"unsure\"\n(0.4–0.6)"],
                [1 - frac, frac], color=[M_ORANGE, M_GRAY], width=0.5)
    axes[1].text(0, 1 - frac + 0.02, f"{1 - frac:.1%}", ha="center", fontsize=11, fontweight="bold")
    axes[1].text(1, frac + 0.02, f"{frac:.1%}", ha="center", fontsize=11)
    axes[1].set_ylim(0, 1.1)
    axes[1].set_ylabel("share of predictions")
    axes[1].set_title("confidence of predictions", fontsize=11)
    for ax in axes:
        ax.spines[["top", "right"]].set_visible(False)
    return save_fig(fig, "calibration.png")


def chart_geometry():
    cats = ["same\nlabel", "different\nlabel", "same\npaper", "different\npaper"]
    keys = ("mean_cosine_same_label", "mean_cosine_diff_label", "mean_cosine_same_paper", "mean_cosine_diff_paper")
    claim = [SIM["claim_vec"][k] for k in keys]
    image = [SIM["image_vec"][k] for k in keys]
    x = np.arange(len(cats))
    w = 0.35
    fig, ax = plt.subplots(figsize=(9, 4.0))
    ax.bar(x - w / 2, claim, w, label="claim embeddings", color=M_NAVY)
    ax.bar(x + w / 2, image, w, label="image embeddings", color=M_ORANGE)
    for i, (c, im) in enumerate(zip(claim, image)):
        ax.text(i - w / 2, c + 0.012, f"{c:.4f}", ha="center", fontsize=8.5)
        ax.text(i + w / 2, im + 0.012, f"{im:.4f}", ha="center", fontsize=8.5)
    ax.set_xticks(x)
    ax.set_xticklabels(cats)
    ax.set_ylabel("mean cosine similarity")
    ax.legend(fontsize=9)
    ax.spines[["top", "right"]].set_visible(False)
    return save_fig(fig, "geometry.png")


# ------------------------------------------------------------------ deck

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = Inches(13.333), Inches(7.5)


def rect(slide, x, y, w, h, color):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def textbox(slide, x, y, w, h, lines, size=16, color=DARK, bullet=True):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(lines):
        text, lvl = item if isinstance(item, tuple) else (item, 0)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ("" if not bullet else ("• " if lvl == 0 else "– ")) + text
        p.level = lvl
        p.font.size = Pt(size if lvl == 0 else size - 2)
        p.font.color.rgb = color
        p.space_after = Pt(8)
    return box


def header(slide, title, tag=""):
    rect(slide, 0, 0, SW, Inches(1.05), NAVY)
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.18), Inches(10.6), Inches(0.75))
    p = box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = WHITE
    if tag:
        tb = slide.shapes.add_textbox(Inches(11.0), Inches(0.3), Inches(2.1), Inches(0.5))
        p = tb.text_frame.paragraphs[0]
        p.text = tag
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = ORANGE


def why_box(slide, text, y=Inches(6.35)):
    x, w = Inches(0.5), Inches(12.3)
    rect(slide, x, y, w, Inches(0.95), GRAY_BG)
    rect(slide, x, y, Inches(0.08), Inches(0.95), ORANGE)
    box = slide.shapes.add_textbox(x + Inches(0.25), y + Inches(0.07), w - Inches(0.5), Inches(0.82))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Why: " + text
    p.font.size = Pt(13)
    p.font.color.rgb = DARK


def divider(number, title, subtitle):
    slide = prs.slides.add_slide(BLANK)
    rect(slide, 0, 0, SW, SH, NAVY)
    box = slide.shapes.add_textbox(Inches(0.9), Inches(2.1), Inches(11.5), Inches(3))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = number
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = ORANGE
    p2 = tf.add_paragraph()
    p2.text = title
    p2.font.size = Pt(40)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p3 = tf.add_paragraph()
    p3.text = subtitle
    p3.font.size = Pt(18)
    p3.font.color.rgb = RGBColor(0xC9, 0xD4, 0xE0)
    return slide


def content(title, tag="Notebook 04"):
    slide = prs.slides.add_slide(BLANK)
    header(slide, title, tag)
    return slide


# 1 — title
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, SW, SH, NAVY)
rect(s, 0, Inches(4.9), SW, Inches(0.06), ORANGE)
box = s.shapes.add_textbox(Inches(0.9), Inches(2.0), Inches(11.5), Inches(2.8))
tf = box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Was 55% Real?"
p.font.size = Pt(44)
p.font.bold = True
p.font.color.rgb = WHITE
for sub, sz in [("A guided tour of notebook 4: the statistical audit of the SciVer baseline", 20),
                ("Teaching edition — every check explained: what it tests, why it was run, what it found", 16)]:
    p = tf.add_paragraph()
    p.text = sub
    p.font.size = Pt(sz)
    p.font.color.rgb = RGBColor(0xC9, 0xD4, 0xE0)

# 2 — the plan of attack
s = content("The question, and the plan of attack")
textbox(s, Inches(0.6), Inches(1.35), Inches(12.1), Inches(4.8), [
    "Notebook 2 reported ~55% accuracy. Guessing gives 50%. The whole audit answers one question: is the gap skill or luck?",
    "The audit proceeds in five moves, each closing a different escape route:",
    ("Replicate — rerun the exact headline number before investigating it.", 1),
    ("Inspect the inputs — what did the model actually see?", 1),
    ("Calibrate chance — how high can a know-nothing model score on this test?", 1),
    ("Probe and stress-test — which input carries the edge, and does it survive resampling?", 1),
    ("Audit the test itself — leakage, selection effects, calibration, and the geometry of the features.", 1),
])
why_box(s, "a single suspicious number is never refuted by one argument. The audit's strength is that several "
           "independent checks — resampling, rebuilding, geometry — all reach the same verdict.")

# 3 — vocabulary
s = content("Vocabulary for the audit")
textbox(s, Inches(0.6), Inches(1.3), Inches(12.1), Inches(5.5), [
    "Null distribution: the range of scores produced by models that cannot know the answer. Anything inside this range is explainable by luck.",
    "Permutation test: build that null empirically — scramble the training labels, retrain, score on the real test set, repeat 300×. The model keeps every capacity it had; only the truth is removed.",
    "p-value: the fraction of scrambled-label models that scored at least as well as the real one. Small p = hard to explain by luck. Two-sided p also counts models as far below 50% (under-performing is suspicious too).",
    "Cross-validation (CV): re-divide the data into many train/test splits and average. A real ability survives re-splitting; a lucky draw does not.",
    "Bootstrap confidence interval: resample the test predictions to see how much the score wobbles on this finite test set.",
    "Family-wise (best-of-N) correction: if you try six inputs and report the best, compare it against the best of six know-nothing models, not one.",
    "Calibration: whether a model's stated probabilities match its actual hit rate.",
], size=14)

# ---- Part 1
divider("PART 1 · SECTIONS 2–3", "Replicate, then inspect the inputs",
        "Never audit a number you haven't reproduced; never trust a score before checking what the model saw")

s = content("Step 1 — Replicate the headline")
textbox(s, Inches(0.6), Inches(1.4), Inches(12.1), Inches(4.4), [
    f"The notebook refits the exact notebook-2 pipeline (scaler + balanced logistic regression, all 2,432 features) on the official split: accuracy {ABL['full_notebook_stack']['accuracy']:.3f}.",
    "In the original environment this same pipeline gave 0.565. The 2-point difference came from a software substitution that resized a few images microscopically differently.",
    ("A result that moves 2 points under numerically irrelevant preprocessing is already announcing that it is noise-dominated.", 1),
])
why_box(s, "replication comes first because every later comparison hangs off this number. Catching the 0.565 → 0.547 "
           "drift early turned a potential confusion into a finding: the headline is preprocessing-noise-dominated.")

s = content("Step 2 — What did the model actually see?")
textbox(s, Inches(0.6), Inches(1.4), Inches(12.1), Inches(4.4), [
    "The evidence-text embedding is checked directly: across all 416 pairs it has one unique row and per-dimension variation of ~1e-6 — it is the embedding of an empty string, identical everywhere.",
    "Confirmed against the raw files: the val/test records carry no caption, OCR, or context fields at all (those live in separate per-paper JSONs the parser never opened).",
    "Consequence: the two claim-vs-evidence interaction blocks compare the claim against a constant — 768 of the 2,432 features are claim-only transforms in disguise, and the model only ever saw the claim and the raw image.",
])
why_box(s, "a claim sentence alone cannot honestly reveal whether a chart supports it. Establishing this up front "
           "defines the audit's expectation: any above-chance text accuracy must be an artifact, not understanding.")

# ---- Part 2
divider("PART 2 · SECTIONS 4–5", "Calibrating chance",
        "Before judging 55%, measure what know-nothing models score on this exact test")

s = content("What pure chance looks like on 276 questions")
s.shapes.add_picture(str(chart_chance()), Inches(1.6), Inches(1.35), width=Inches(10.2))
why_box(s, "with 276 questions, guessing has a standard deviation of 3 points — the 95% chance band reaches 55.9%, "
           "and reliably detecting a real effect needs ~58.4% true accuracy. The headline sits inside the band. "
           "Eyeballing \"55 > 50\" was never enough.", y=Inches(6.15))

s = content("The permutation test: measure the null, don't assume it")
s.shapes.add_picture(str(chart_null_ranges()), Inches(1.6), Inches(1.3), width=Inches(10.2))
why_box(s, "scrambled-label models keep every capacity of the real one (2,432 features, refitting, the same test set) "
           "and still span 37%–60% across 300 runs, symmetrically around 50%. This empirical null prices in everything "
           "the textbook formula misses. Both tails are recorded; under-performance would be flagged too.", y=Inches(6.15))

# ---- Part 3
divider("PART 3 · SECTIONS 5–7", "The probes",
        "Isolate each input, then apply the decisive test: does the edge survive resampling?")

s = content("Probe 1 — Can the claim's wording alone give the answer away?")
textbox(s, Inches(0.6), Inches(1.4), Inches(12.1), Inches(4.4), [
    "Refuted claims are edited copies of true sentences (notebook 3's perturbation fields). Editing can leave a stylistic fingerprint a text model could detect — scoring without ever checking the chart.",
    "Test: TF-IDF logistic regression on the claim string only. Since the claim cannot honestly decide entailment, this probe is a pure artifact detector.",
    (f"Main subset (n=276): {TF['direct_chart']['accuracy']:.3f}, p = {TF['direct_chart']['perm_p']:.2f} — no fingerprint detectable.", 1),
    (f"Full set (n=996): {TF['all_chart_table']['accuracy']:.3f}, p = {TF['all_chart_table']['perm_p']:.2f} (two-sided {TF['all_chart_table']['perm_p_two_sided']:.3f}) — a faint, real fingerprint, visible only because the test set is large.", 1),
])
why_box(s, "this probe separates two explanations for any text contribution: understanding (impossible by construction) "
           "vs dataset-authoring artifact. The faint large-n fingerprint is a caution for anyone training text models on SciVer.")

s = content("Probe 2 + the decisive check: single split vs resampling")
s.shapes.add_picture(str(chart_ablation_cv()), Inches(1.6), Inches(1.3), width=Inches(10.2))
why_box(s, "on the one official split, image-driven blocks look interesting (0.565–0.576). Under 50 resplits — with "
           "2.4× more training data — every edge collapses to chance. A real ability improves with more data; "
           "a lucky train/test draw evaporates. This is the single most decisive chart in the project.", y=Inches(6.15))

# ---- Part 4
divider("PART 4 · SECTIONS 8–9", "Could the test itself be rigged?",
        "Leakage, selection effects, and grouping — three ways a fair-looking test can lie")

s = content("Integrity checks: leakage and grouping")
textbox(s, Inches(0.6), Inches(1.4), Inches(12.1), Inches(4.4), [
    "Contamination: zero images and zero papers appear on both sides of the official split — the single-split number cannot be memorization. The direct+chart subset has zero duplicated images, so CV folds were clean too.",
    "Duplication elsewhere: 48 images in the analytical/table subsets carry both labels (a deterministic 0.968 accuracy ceiling for anything image-aware) — irrelevant here, critical for future all-pairs work.",
    "Grouped CV: re-running cross-validation with each paper confined to one side of every fold changed nothing (e.g. image block 0.519 grouped vs 0.519 plain).",
])
why_box(s, "before concluding \"lucky split\", its strongest rival explanations — memorization and paper-level leakage — "
           "had to be tested and excluded. They were.")

s = content("Selection: we tried six inputs and kept the best")
s.shapes.add_picture(str(chart_six_mirror()), Inches(1.6), Inches(1.35), width=Inches(10.2))
why_box(s, f"the best of six know-nothing models averages {FW['null_max']['mean']:.1%} — selection alone inflates scores. "
           f"The observed best ({FW['observed_best_accuracy']:.1%}) is still unusual under this harsher null "
           f"(family-wise p = {FW['family_wise_perm_p']}), so selection alone does not explain it — "
           "but the CV collapse already did. The mirror-image worst-of-six null (lower-tail p = "
           f"{FW['family_wise_perm_p_low']}) confirms the test is symmetric and nothing under-performs suspiciously.",
        y=Inches(6.0))

# ---- Part 5
divider("PART 5 · SECTIONS 11–13", "Root cause",
        "More data, tuning, probabilities, and finally the geometry of the embeddings")

s = content("Scaling up: 996 test examples, same verdict — sharper")
s.shapes.add_picture(str(chart_bigpop()), Inches(1.8), Inches(1.3), width=Inches(9.7))
why_box(s, "at n=996 the test can detect ~4-point effects — and finds none: the full stack lands at 0.498, dead chance. "
           "Metadata one-hots score 0.490 (the image whisper is not a chart-vs-table base-rate artifact), and honest "
           "regularization tuning peaks at ~0.51 CV accuracy. More data sharpened the verdict instead of rescuing the model.",
        y=Inches(6.05))

s = content("Calibration: confidently wrong is worse than unsure")
s.shapes.add_picture(str(chart_calibration()), Inches(1.8), Inches(1.35), width=Inches(9.7))
why_box(s, f"a useful model says \"60% sure\" and is right 60% of the time. This one makes near-certain predictions "
           f"{1 - CAL['frac_predictions_in_0.4_0.6']:.0%} of the time while being right about half the time — its Brier score "
           f"({CAL['brier_score']}) is worse than always answering \"no idea\" ({CAL['brier_score_always_0.5']}). With 2,432 "
           "features and ~500 training rows, logistic regression memorizes and exports confidence, not information.",
        y=Inches(6.05))

s = content("The root cause: the label is not in the numbers")
s.shapes.add_picture(str(chart_geometry()), Inches(1.8), Inches(1.3), width=Inches(9.7))
why_box(s, "entailed and refuted examples are equidistant to the fourth decimal — no classifier can separate classes "
           "the geometry does not encode. What is encoded is which paper an example came from (0.776 vs 0.663 for images), "
           f"and labels cluster within papers (p = {PLB['perm_p']}): a lurking shortcut that grouped evaluation must always block.",
        y=Inches(6.05))

# closing
s = content("The verdict, and what made it trustworthy")
textbox(s, Inches(0.6), Inches(1.35), Inches(12.1), Inches(4.9), [
    "The ~55% was a favorable train/test draw, not skill. Three independent lines of evidence agree:",
    ("Resampling: every edge collapses under repeated and grouped CV despite more training data.", 1),
    ("Rebuilding: the headline moved 0.565 → 0.547 under an irrelevant preprocessing change.", 1),
    ("Geometry: the embeddings do not encode the label at all — there was nothing to find.", 1),
    "The checks that came back clean matter as much as the ones that didn't: no leakage, selection priced in, both tails of every null recorded.",
    "What happens next — recovering the unused captions, sizing a properly-powered retry, and the untouched multi-visual half of SciVer — is notebook 05.",
])
why_box(s, "a negative result is a deliverable when it is this well-documented: every number traces to a committed "
           "results file, and every alternative explanation was tested rather than assumed.")

prs.save(OUT_PPTX)
print(f"Wrote {OUT_PPTX} ({OUT_PPTX.stat().st_size / 1024:.0f} KiB, {len(prs.slides._sldIdLst)} slides)")
