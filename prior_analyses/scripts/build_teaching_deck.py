"""Build a styled teaching deck (PowerPoint) covering notebooks 01-03.

Generates docs/notebooks_1_3_teaching_deck.pptx. The deck explains, in
classroom style, what notebooks 01 (Qdrant feature store), 02 (logistic-
regression baseline), and 03 (raw-to-processed mapping) do and why.

Number sources:
  - Dataset/pipeline counts: re-read live from
    data/processed/sciver_pair_features.parquet (same file the notebooks use).
  - Model results (accuracy, confusion matrix, coefficient norms): recomputed
    live with exactly the notebook-02 pipeline, so they match the committed
    notebook outputs (accuracy 0.547 in this environment).
  - Nearest-neighbor cosine scores: committed outputs of notebook 01.

Usage:
    python scripts/build_teaching_deck.py
"""

from __future__ import annotations

import tempfile
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import accuracy_score, balanced_accuracy_score, confusion_matrix
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import StandardScaler

REPO_ROOT = Path(__file__).resolve().parents[1]
OUT_PPTX = REPO_ROOT / "docs/notebooks_1_3_teaching_deck.pptx"
CHART_DIR = Path(tempfile.mkdtemp(prefix="teach_charts_"))

# Palette
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
ORANGE = RGBColor(0xE0, 0x7B, 0x39)
DARK = RGBColor(0x26, 0x26, 0x26)
GRAY_BG = RGBColor(0xF0, 0xF2, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
M_NAVY, M_ORANGE, M_GRAY = "#1F3A5F", "#E07B39", "#9AA5B1"

# ------------------------------------------------------------- live numbers

df = pd.read_parquet(REPO_ROOT / "data/processed/sciver_pair_features.parquet")
VEC_COLS = ["claim_vec", "evidence_text_vec", "pair_text_vec", "image_vec"]
for c in VEC_COLS:
    df = df[df[c].notna()]
N_PAIRS = len(df)                                   # 1,500
N_RAW_EXAMPLES = 3000                               # notebook 03: parsed val+test examples
mod_counts = df["modality"].value_counts().to_dict()
ct_counts = df["claim_type"].value_counts().to_dict()

dc = df[(df["modality"] == "chart") & (df["claim_type"] == "direct")].reset_index(drop=True)
split_label = dc.groupby(["split", "label"]).size()

def stack(frame, col):
    return np.vstack([np.asarray(v, dtype=np.float32) for v in frame[col]])

claim_v, ev_v = stack(dc, "claim_vec"), stack(dc, "evidence_text_vec")
pair_v, img_v = stack(dc, "pair_text_vec"), stack(dc, "image_vec")
blocks = {
    "claim_vec": claim_v,
    "evidence_text_vec": ev_v,
    "pair_text_vec": pair_v,
    "image_vec": img_v,
    "abs(claim − evidence)": np.abs(claim_v - ev_v),
    "claim × evidence": claim_v * ev_v,
}
X = np.hstack(list(blocks.values())).astype(np.float64)
y = dc["label_id"].astype(int).to_numpy()
split = dc["split"].astype(str).to_numpy()
tr, te = split == "val", split == "test"

model = Pipeline([
    ("scaler", StandardScaler()),
    ("classifier", LogisticRegression(max_iter=2000, class_weight="balanced", random_state=42)),
]).fit(X[tr], y[tr])
pred = model.predict(X[te])
ACC = accuracy_score(y[te], pred)
BACC = balanced_accuracy_score(y[te], pred)
CM = confusion_matrix(y[te], pred, labels=[0, 1])

coef = model.named_steps["classifier"].coef_.ravel()
coef_norms, start = {}, 0
for name, b in blocks.items():
    coef_norms[name] = float(np.linalg.norm(coef[start:start + b.shape[1]]))
    start += b.shape[1]

# Committed notebook-01 outputs: top non-self neighbor cosine per vector space.
NN_TOP1 = {"pair_text_vec": 0.682, "claim_vec": 0.665, "image_vec": 0.843}

# ------------------------------------------------------------------ charts

def save_fig(fig, name):
    p = CHART_DIR / name
    fig.savefig(p, dpi=200, bbox_inches="tight")
    plt.close(fig)
    return p


def chart_funnel():
    n_train, n_test = int(tr.sum()), int(te.sum())
    stages = [
        ("raw examples parsed (val + test)", N_RAW_EXAMPLES, M_GRAY),
        ("usable pairs: exactly one chart/table per claim", N_PAIRS, M_NAVY),
        ("this baseline: chart images + direct claims", len(dc), M_NAVY),
    ]
    fig, ax = plt.subplots(figsize=(9, 4))
    ys = [3, 2, 1]
    for (name, v, color), ypos in zip(stages, ys):
        ax.barh(ypos, v, color=color, height=0.6)
        ax.text(v + 40, ypos, f"{v:,}", va="center", fontsize=11, fontweight="bold", color="#262626")
    # Final stage: the same 416 divided into train and test, not filtered further.
    ax.barh(0, n_train, color=M_NAVY, height=0.6)
    ax.barh(0, n_test, left=n_train, color=M_ORANGE, height=0.6)
    ax.text(n_train / 2, 0, f"train\n{n_train}", va="center", ha="center", fontsize=10,
            fontweight="bold", color="white")
    ax.text(n_train + n_test / 2, 0, f"test\n{n_test}", va="center", ha="center", fontsize=10,
            fontweight="bold", color="white")
    ax.set_yticks(ys + [0])
    ax.set_yticklabels([s[0] for s in stages] + ["the same 416, divided for evaluation"], fontsize=10)
    ax.set_xlim(0, N_RAW_EXAMPLES * 1.12)
    ax.set_xlabel("number of examples")
    ax.spines[["top", "right"]].set_visible(False)
    return save_fig(fig, "funnel.png")


def chart_neighbors():
    fig, ax = plt.subplots(figsize=(7.5, 3.4))
    names = ["claim text\n(MiniLM, 384-d)", "claim + evidence text\n(MiniLM, 384-d)", "chart image\n(CLIP, 512-d)"]
    vals = [NN_TOP1["claim_vec"], NN_TOP1["pair_text_vec"], NN_TOP1["image_vec"]]
    bars = ax.bar(names, vals, color=[M_NAVY, M_NAVY, M_ORANGE], width=0.55)
    for b, v in zip(bars, vals):
        ax.text(b.get_x() + b.get_width() / 2, v + 0.015, f"{v:.2f}", ha="center", fontsize=11, fontweight="bold")
    ax.set_ylim(0, 1.0)
    ax.set_ylabel("cosine similarity of closest neighbor")
    ax.spines[["top", "right"]].set_visible(False)
    return save_fig(fig, "neighbors.png")


def chart_label_balance():
    fig, ax = plt.subplots(figsize=(7.5, 3.6))
    groups = [("val (train)", split_label.get(("val", "entailed"), 0), split_label.get(("val", "refuted"), 0)),
              ("test", split_label.get(("test", "entailed"), 0), split_label.get(("test", "refuted"), 0))]
    x = np.arange(len(groups))
    w = 0.35
    ax.bar(x - w / 2, [g[1] for g in groups], w, label="entailed", color=M_NAVY)
    ax.bar(x + w / 2, [g[2] for g in groups], w, label="refuted", color=M_ORANGE)
    for i, g in enumerate(groups):
        ax.text(i - w / 2, g[1] + 2, str(g[1]), ha="center", fontsize=11)
        ax.text(i + w / 2, g[2] + 2, str(g[2]), ha="center", fontsize=11)
    ax.set_xticks(x)
    ax.set_xticklabels([g[0] for g in groups])
    ax.set_ylabel("examples")
    ax.legend(frameon=False)
    ax.spines[["top", "right"]].set_visible(False)
    return save_fig(fig, "label_balance.png")


def chart_feature_blocks():
    fig, ax = plt.subplots(figsize=(9, 3.6))
    names = list(blocks)
    dims = [b.shape[1] for b in blocks.values()]
    colors = [M_NAVY, M_GRAY, M_NAVY, M_ORANGE, M_GRAY, M_GRAY]
    bars = ax.bar(names, dims, color=colors)
    for b, d in zip(bars, dims):
        ax.text(b.get_x() + b.get_width() / 2, d + 8, str(d), ha="center", fontsize=11)
    ax.set_ylabel("number of features")
    ax.set_title(f"total: {X.shape[1]:,} features per example — learned from only {int(tr.sum())} training examples",
                 fontsize=11)
    plt.xticks(rotation=12, ha="right", fontsize=9)
    ax.spines[["top", "right"]].set_visible(False)
    return save_fig(fig, "feature_blocks.png")


def chart_confusion():
    fig, ax = plt.subplots(figsize=(5.2, 4.2))
    im = ax.imshow(CM, cmap="Blues", vmin=0)
    for i in range(2):
        for j in range(2):
            ax.text(j, i, str(CM[i, j]), ha="center", va="center", fontsize=18,
                    color="white" if CM[i, j] > CM.max() * 0.6 else "#262626", fontweight="bold")
    ax.set_xticks([0, 1]); ax.set_yticks([0, 1])
    ax.set_xticklabels(["predicted refuted", "predicted entailed"], fontsize=10)
    ax.set_yticklabels(["true refuted", "true entailed"], fontsize=10)
    ax.set_title(f"accuracy {ACC:.3f} — balanced accuracy {BACC:.3f}", fontsize=11)
    return save_fig(fig, "confusion.png")


def chart_coef_norms():
    fig, ax = plt.subplots(figsize=(9, 3.6))
    names = list(coef_norms)
    vals = list(coef_norms.values())
    colors = [M_ORANGE if v < 1e-9 else M_NAVY for v in vals]
    bars = ax.barh(names[::-1], vals[::-1], color=colors[::-1])
    for b, v in zip(bars, vals[::-1]):
        ax.text(b.get_width() + 0.012, b.get_y() + b.get_height() / 2,
                f"{v:.3f}" + ("  ← exactly zero: the input is constant!" if v < 1e-9 else ""),
                va="center", fontsize=10)
    ax.set_xlim(0, max(vals) * 1.55)
    ax.set_xlabel("size of learned weights per feature block (L2 norm)")
    ax.spines[["top", "right"]].set_visible(False)
    return save_fig(fig, "coef_norms.png")


# ------------------------------------------------------------------ deck

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = Inches(13.333), Inches(7.5)


def rect(slide, x, y, w, h, color, line=False):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    if not line:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def textbox(slide, x, y, w, h, lines, size=16, color=DARK, bold_first=False, bullet=True):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(lines):
        text, lvl = item if isinstance(item, tuple) else (item, 0)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        prefix = "" if not bullet else ("• " if lvl == 0 else "– ")
        p.text = prefix + text
        p.font.size = Pt(size if lvl == 0 else size - 2)
        p.font.color.rgb = color
        p.font.bold = bold_first and i == 0
        p.space_after = Pt(8)
    return box


def header(slide, title, tag=""):
    rect(slide, 0, 0, SW, Inches(1.05), NAVY)
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.18), Inches(10.6), Inches(0.75))
    p = box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = WHITE
    if tag:
        tb = slide.shapes.add_textbox(Inches(11.0), Inches(0.3), Inches(2.1), Inches(0.5))
        p = tb.text_frame.paragraphs[0]
        p.text = tag
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = ORANGE


def why_box(slide, text, y=Inches(6.35), x=Inches(0.5), w=None):
    w = w or Inches(12.3)
    rect(slide, x, y, w, Inches(0.85), GRAY_BG)
    rect(slide, x, y, Inches(0.08), Inches(0.85), ORANGE)
    box = slide.shapes.add_textbox(x + Inches(0.25), y + Inches(0.08), w - Inches(0.5), Inches(0.7))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Why: " + text
    p.font.size = Pt(14)
    p.font.color.rgb = DARK


def divider(number, title, subtitle):
    slide = prs.slides.add_slide(BLANK)
    rect(slide, 0, 0, SW, SH, NAVY)
    box = slide.shapes.add_textbox(Inches(0.9), Inches(2.1), Inches(11.5), Inches(3))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = number
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = ORANGE
    p2 = tf.add_paragraph()
    p2.text = title
    p2.font.size = Pt(40)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p3 = tf.add_paragraph()
    p3.text = subtitle
    p3.font.size = Pt(18)
    p3.font.color.rgb = RGBColor(0xC9, 0xD4, 0xE0)
    return slide


def content(title, tag=""):
    slide = prs.slides.add_slide(BLANK)
    header(slide, title, tag)
    return slide


# 1 — title
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, SW, SH, NAVY)
rect(s, 0, Inches(4.9), SW, Inches(0.06), ORANGE)
box = s.shapes.add_textbox(Inches(0.9), Inches(2.0), Inches(11.5), Inches(2.8))
tf = box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "From Raw Dataset to First Baseline"
p.font.size = Pt(44)
p.font.bold = True
p.font.color.rgb = WHITE
for sub, sz in [("A guided tour of notebooks 1–3: the SciVer data, the vector database, and the first model", 20),
                ("Teaching edition — every step explained: what was done, and why", 16)]:
    p = tf.add_paragraph()
    p.text = sub
    p.font.size = Pt(sz)
    p.font.color.rgb = RGBColor(0xC9, 0xD4, 0xE0)

# 2 — course map
s = content("The three notebooks, in the order the data flows")
textbox(s, Inches(0.6), Inches(1.4), Inches(12.1), Inches(4.6), [
    "Notebook 03 — Raw → Processed mapping. Where the data comes from: the SciVer dataset files, how 3,000 raw rows become 1,500 clean claim-chart pairs, and how every record gets a traceable ID.",
    "Notebook 01 — The vector database. Where the data lives: each pair is stored with four embedding vectors and its metadata, ready to query, filter, and reuse across experiments.",
    "Notebook 02 — The first model. What we do with it: train logistic regression on the embeddings to predict entailed vs refuted, and get the ~55% number the rest of the project investigates.",
    "We follow that order (3 → 1 → 2) because it matches the pipeline: parse the data, store the features, fit a model.",
])
why_box(s, "separating these stages means any experiment can be traced back: model → features → raw file. "
           "When a result looks wrong (and one did), each stage can be audited on its own.")

# 3 — vocabulary
s = content("Vocabulary for the rest of the deck")
textbox(s, Inches(0.6), Inches(1.35), Inches(12.1), Inches(5.4), [
    "Claim: one sentence asserting something about a figure (e.g. \"AUC increases by 3.98 points…\").",
    "Label: entailed (the figure supports the claim) or refuted (it contradicts it). This is what we predict.",
    "Embedding / vector: a fixed-length list of numbers representing a sentence or image, produced by a pretrained model. Similar content gets similar vectors.",
    "Cosine similarity: a score in [−1, 1] for how similar two vectors point. 1 = identical direction.",
    "Vector database (Qdrant): a store that holds vectors plus metadata and can return the nearest vectors to a query.",
    "Train / test split: the model learns from training examples and is graded on held-out test examples it never saw.",
    "Logistic regression: the simplest standard classifier — a weighted sum of the features pushed through a probability curve.",
    "Label leakage: any path by which the answer sneaks into the inputs. Preventing it is a design requirement, not an afterthought.",
], size=15)

# ---- Part 1: notebook 03
divider("PART 1 · NOTEBOOK 03", "Where the data comes from",
        "Raw SciVer files → parsed examples → supervised pairs, with provenance at every step")

s = content("The raw material: SciVer", "Notebook 03")
textbox(s, Inches(0.6), Inches(1.4), Inches(12.1), Inches(4.6), [
    "SciVer is a published benchmark: scientific claims paired with real figures and tables from arXiv papers.",
    "Two JSON files (valset.json, testset.json) hold 3,000 rows. Each row: a claim, a pointer to an image file, a label, and a claim type (direct = reads a value off the figure; analytical = requires reasoning).",
    "A download manifest records which dataset revision was fetched — so anyone can confirm they have the same data.",
    ("Foreshadowing: the parsed example shown in the notebook has an empty caption field. That detail becomes the project's central finding later.", 1),
])
why_box(s, "every analysis is only as good as its input files. Recording exactly what was downloaded (and looking at "
           "one raw row with your own eyes) is the cheapest insurance in the pipeline.")

s = content("From 3,000 rows to one clean supervised unit", "Notebook 03")
s.shapes.add_picture(str(chart_funnel()), Inches(1.4), Inches(1.35), width=Inches(10.5))
why_box(s, "each kept pair has exactly one figure per claim — a clean (input, answer) unit for supervised learning. "
           f"The {N_PAIRS:,} pairs stay balanced: {mod_counts.get('chart', 0)} chart / {mod_counts.get('table', 0)} table, "
           f"{ct_counts.get('direct', 0)} direct / {ct_counts.get('analytical', 0)} analytical, ~50/50 entailed/refuted.")

s = content("Provenance: deterministic IDs and leakage control", "Notebook 03")
textbox(s, Inches(0.6), Inches(1.4), Inches(12.1), Inches(4.6), [
    "Every claim, evidence item, and pair gets an ID computed from its content (a deterministic UUID). Re-running the pipeline reproduces identical IDs — records can be matched across runs and tools.",
    "The raw rows contain the answer key: fields like origin_statement, perturbed_statement, and perturbed_explanation describe how refuted claims were made. The pipeline embeds only the claim field and keeps the answer-key fields out of all features.",
    "Honest lesson from the notebook itself: its demo that rebuilds pairs in memory and joins them to the stored manifest matched 0 of 1,500 rows, because the rebuild used different parser options than the build script.",
    ("Lesson: a determinism claim is only verified when the check replicates the original settings exactly.", 1),
])
why_box(s, "if the answer-key fields leaked into the features, the model could score high while learning nothing about "
           "charts. Keeping them out is what makes any later accuracy worth discussing.")

# ---- Part 2: notebook 01
divider("PART 2 · NOTEBOOK 01", "The vector database",
        "1,500 pairs stored as vectors + metadata in Qdrant, ready to query and reuse")

s = content("What one stored record looks like", "Notebook 01")
textbox(s, Inches(0.6), Inches(1.35), Inches(12.1), Inches(5.0), [
    "Each of the 1,500 pairs is one point in the database, holding two things:",
    ("Four named vectors — claim_vec (384 numbers, MiniLM text model), evidence_text_vec (384, same model), "
     "pair_text_vec (384, claim + evidence together), image_vec (512, CLIP image model). Compared by cosine similarity.", 1),
    ("A payload of metadata — the claim text, label, split (val/test), claim type, modality, paper ID, image path, "
     "and which embedding models produced the vectors.", 1),
    "Three ways to get records out:",
    ("Scroll: browse records in storage order (no ranking) — for inspection.", 1),
    ("Nearest-neighbor search: give a vector, get the most similar points — optionally restricted by metadata filters "
     "(e.g. only val + entailed + chart).", 1),
    ("Retrieve by ID: the deterministic pair ID maps to the database ID, so an exact record can be fetched directly.", 1),
], size=15)
why_box(s, "embedding 1,500 texts and images is slow; storing the vectors once lets every later experiment "
           "(notebooks 02 and 04, all the scripts) reuse identical features. Metadata filters make split-aware "
           "retrieval possible, which prevents accidental train/test mixing.")

s = content("What the embeddings see: a first clue", "Notebook 01")
s.shapes.add_picture(str(chart_neighbors()), Inches(2.4), Inches(1.45), width=Inches(8.5))
why_box(s, "in the notebook's example search, the closest other image scores 0.84 while the closest other text scores "
           "~0.67. Chart images crowd together in CLIP space — scientific figures look alike to a general-purpose image "
           "model. Notebook 04 later shows this crowding encodes which paper a figure came from, not whether the claim is true.",
        y=Inches(5.9))

# ---- Part 3: notebook 02
divider("PART 3 · NOTEBOOK 02", "The first model",
        "Logistic regression on the stored embeddings: chart + claim → entailed or refuted?")

s = content("The task and the split", "Notebook 02")
s.shapes.add_picture(str(chart_label_balance()), Inches(3.0), Inches(1.4), width=Inches(7.3))
why_box(s, f"we restrict to direct claims about charts ({len(dc)} pairs): the value is supposed to be readable off the figure, "
           "so it is the fairest subset for a simple model. We train on the official val split (140) and grade on the "
           "official test split (276) — examples the model has never seen. Both sides are near 50/50, so guessing scores ~50%.",
        y=Inches(5.75))

s = content("The feature matrix: 2,432 numbers per example", "Notebook 02")
s.shapes.add_picture(str(chart_feature_blocks()), Inches(1.7), Inches(1.4), width=Inches(10.0))
why_box(s, "the two interaction blocks (difference and product of claim and evidence vectors) are meant to measure "
           "claim-vs-evidence agreement. With far more features (2,432) than training examples (140), the model can "
           "memorize the training set — one reason test-set grading and the later statistics matter.",
        y=Inches(5.85))

s = content("The model: scale, then draw one line", "Notebook 02")
textbox(s, Inches(0.6), Inches(1.4), Inches(12.1), Inches(4.4), [
    "Step 1 — StandardScaler: rescale every feature to mean 0, variance 1. Text and image vectors have different "
    "numeric ranges; without scaling, big-range features dominate by accident rather than by usefulness.",
    "Step 2 — LogisticRegression(class_weight=\"balanced\"): learn one weight per feature; the weighted sum gives the "
    "probability the pair is entailed. Class weighting compensates for the train split's mild 76/64 imbalance.",
    "Why such a simple model? It is fast, deterministic, and well understood. If real signal exists in the embeddings, "
    "a linear probe usually finds at least part of it; if a linear probe finds nothing, the burden of proof rises sharply.",
])
why_box(s, "baselines come before fancy models: they establish whether the features contain accessible signal at all, "
           "and they are cheap enough to audit thoroughly — which is exactly what happened next.")

s = content(f"The result: {ACC:.1%} accuracy", "Notebook 02")
s.shapes.add_picture(str(chart_confusion()), Inches(0.9), Inches(1.45), width=Inches(5.6))
textbox(s, Inches(7.0), Inches(1.7), Inches(5.8), Inches(4.4), [
    f"Accuracy {ACC:.3f}, balanced accuracy {BACC:.3f}.",
    f"The model gets {CM[0,0]} of {CM[0].sum()} refuted and {CM[1,1]} of {CM[1].sum()} entailed test pairs right — "
    "barely better than alternating coin flips.",
    "Is 54.7% \"better than 50%\"? With only 276 test questions, chance alone spans roughly 44%–56%. "
    "That question is the entire subject of notebook 04 and the statistics deck.",
    ("Answer there: no — the edge does not survive cross-validation or a rebuild.", 1),
], size=15)
why_box(s, "a confusion matrix shows where the errors live: here they are spread evenly across both classes — "
           "the signature of a model with no systematic skill.", y=Inches(6.25))

s = content("Red flags visible before any statistics", "Notebook 02")
s.shapes.add_picture(str(chart_coef_norms()), Inches(1.7), Inches(1.35), width=Inches(10.0))
why_box(s, "the model assigns exactly zero weight to evidence_text_vec — possible only if that input never varies. "
           "It was the embedding of an empty string for all 1,500 pairs (the caption was never loaded — the empty "
           "caption from notebook 03!). The model never saw the evidence it was supposed to check claims against.",
        y=Inches(5.8))

# wrap-up
s = content("What notebooks 1–3 leave us with")
textbox(s, Inches(0.6), Inches(1.4), Inches(12.1), Inches(4.6), [
    "A reproducible pipeline: raw SciVer files → parsed examples → deterministic pair records → embeddings in a "
    "queryable vector store → an exported feature table.",
    f"A first measurement: {ACC:.1%} test accuracy from logistic regression on {X.shape[1]:,} embedding features.",
    "Two warnings already in plain sight: the evidence input is constant (weight exactly zero), and chart embeddings "
    "crowd together regardless of label.",
    "The natural next question — is 55% real skill or luck? — is answered in notebook 04 and the companion deck "
    "(plain_language_summary.pptx): it was luck.",
])
why_box(s, "good pipelines do not guarantee good results — they guarantee that results can be interrogated. "
           "These three notebooks are what made the later verdict possible.")

prs.save(OUT_PPTX)
print(f"Wrote {OUT_PPTX} ({OUT_PPTX.stat().st_size / 1024:.0f} KiB, {len(prs.slides._sldIdLst)} slides)")
print(f"Recomputed NB02 result: acc={ACC:.3f}, bacc={BACC:.3f}, cm={CM.tolist()}")
